#!/usr/bin/env python3
"""Parse the `pushback_300` materials (CS 598 HCI, Group 6).

Inputs (NOT in this repo — supply your own copies; they contain third-party conversation text):
  source/pushback_300_contexts.md   full conversation context for 300 candidate pushbacks
  source/pushback_300_coded.xlsm    300 rows of metadata + AI preliminary theme/subtheme codes

Outputs:
  data_local/contexts.jsonl         event_id, messages, target_index, source  (gitignored)
  data_local/meta.jsonl             workbook metadata joined by event_id       (gitignored)
  results/corpus_stats.json         aggregate counts only                      (committed)

Design notes
  - The markdown is the authority for conversation content; the spreadsheet's
    "Previous user turn" and "Challenged assistant answer" cells are truncated excerpts.
  - Nothing here executes the .xlsm; macros are never run.
  - Quotes and conversation text stay in data_local/ and must not be committed.
"""
import argparse
import collections
import csv
import json
import os
import re
import statistics
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
DATA_LOCAL = os.path.join(HERE, "data_local")
RESULTS = os.path.join(HERE, "results")

BLOCK_RE = re.compile(r"^## (\d{3})\.\s+(PB_[0-9a-f]+)\s*$", re.M)
HEAD_RE = re.compile(r"^\*\*Message (\d+): (user|assistant)([^*]*)\*\*\s*$", re.M)
FENCE_RE = re.compile(r"```(?:text)?\n(.*?)\n```", re.S)
SOURCE_RE = re.compile(r"Source:\s*`([^`]+)`")
FLAG_RE = re.compile(r"AI is_pushback:\s*(\w+)", re.I)


def parse_contexts(path):
    """Return one record per case, in file order."""
    text = open(path, encoding="utf-8").read()
    starts = [(m.start(), m.group(1), m.group(2)) for m in BLOCK_RE.finditer(text)]
    recs = []
    for i, (pos, no, eid) in enumerate(starts):
        end = starts[i + 1][0] if i + 1 < len(starts) else len(text)
        block = text[pos:end]
        parts = HEAD_RE.split(block)
        messages = []
        for j in range(1, len(parts), 4):
            num, role, flag, rest = parts[j], parts[j + 1], parts[j + 2], parts[j + 3]
            fence = FENCE_RE.search(rest)
            body = (fence.group(1) if fence else rest).strip()
            messages.append({
                "index": int(num),
                "role": role.strip(),
                "flag": flag.strip().strip("\u2014").strip(),
                "content": body,
            })
        if not messages:
            continue
        src = SOURCE_RE.search(block)
        target = [m for m in messages if "TARGET USER TURN" in m["flag"]]
        prior = [m for m in messages if "PRECEDING ASSISTANT ANSWER" in m["flag"]]
        recs.append({
            "sample_no": int(no),
            "event_id": eid,
            "source": src.group(1) if src else None,
            "is_pushback_ai": (FLAG_RE.search(block).group(1).lower()
                               if FLAG_RE.search(block) else None),
            "messages": messages,
            "target_index": target[0]["index"] if target else None,
            "target_text": target[0]["content"] if target else None,
            "preceding_answer": prior[-1]["content"] if prior else None,
            "n_messages": len(messages),
        })
    return recs


def parse_workbook(path):
    """Metadata + AI preliminary codes, keyed by event_id. Read-only; no macros executed."""
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl is required: pip install openpyxl")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_vba=False)
    ws = wb["Pushback 300"] if "Pushback 300" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    def g(r, i):
        return r[i] if i < len(r) else None

    # locate the header row: it starts with "Sample no."
    start = None
    for i, r in enumerate(rows):
        if str(g(r, 0) or "").strip().lower().startswith("sample"):
            start = i + 1
            break
    if start is None:
        sys.exit("could not find the header row in the workbook")

    meta = {}
    for r in rows[start:]:
        eid = g(r, 1)
        if not eid:
            continue
        themes = [t.strip() for t in re.split(r"[\n|]", str(g(r, 6) or "")) if t.strip()]
        subs = [s.strip() for s in re.split(r"[\n|]", str(g(r, 7) or "")) if s.strip()]
        meta[str(eid).strip()] = {
            "sample_no": g(r, 0),
            "prev_user_excerpt": g(r, 2),
            "answer_excerpt": g(r, 3),
            "is_pushback_ai": str(g(r, 4) or "").strip().lower(),
            "ai_rationale": g(r, 5),
            "themes_ai": themes,
            "subthemes_ai": subs,
            "quote": g(r, 8),
            "team_notes": g(r, 9),
            "source_ref": g(r, 10),
        }
    return meta


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--contexts", default=os.path.join(HERE, "source", "pushback_300_contexts.md"))
    ap.add_argument("--workbook", default=os.path.join(HERE, "source", "pushback_300_coded.xlsm"))
    args = ap.parse_args()

    os.makedirs(DATA_LOCAL, exist_ok=True)
    os.makedirs(RESULTS, exist_ok=True)

    recs = parse_contexts(args.contexts)
    meta = parse_workbook(args.workbook)
    print(f"parsed {len(recs)} cases from markdown, {len(meta)} rows from workbook")

    have_meta = sum(1 for r in recs if r["event_id"] in meta)
    print(f"joined on event_id: {have_meta}/{len(recs)}")
    if have_meta != len(recs):
        missing = [r["event_id"] for r in recs if r["event_id"] not in meta][:5]
        print("  missing metadata for:", missing)

    # structural checks
    bad_target = [r["event_id"] for r in recs
                  if r["target_index"] is None
                  or r["target_index"] != r["messages"][-1]["index"]
                  or r["messages"][-1]["role"] != "user"]
    no_prior = [r["event_id"] for r in recs if not r["preceding_answer"]]
    print(f"target is last user message: {len(recs)-len(bad_target)}/{len(recs)}")
    print(f"has a preceding assistant answer: {len(recs)-len(no_prior)}/{len(recs)}")

    with open(os.path.join(DATA_LOCAL, "contexts.jsonl"), "w") as f:
        for r in recs:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    with open(os.path.join(DATA_LOCAL, "meta.jsonl"), "w") as f:
        for eid, m in meta.items():
            f.write(json.dumps({"event_id": eid, **m}, ensure_ascii=False) + "\n")

    # aggregate statistics only -> safe to commit
    sub_counts = collections.Counter()
    theme_counts = collections.Counter()
    for m in meta.values():
        for s in m["subthemes_ai"]:
            code = s.split()[0].strip("—-").strip()
            if code:
                sub_counts[code] += 1
        for t in m["themes_ai"]:
            theme_counts[t.strip()] += 1

    tgt_len = [len(r["target_text"] or "") for r in recs]
    ans_len = [len(r["preceding_answer"] or "") for r in recs]
    ctx_len = [sum(len(m["content"]) for m in r["messages"][:-1]) for r in recs]
    stats = {
        "n_cases": len(recs),
        "n_joined": have_meta,
        "n_source_shards": len({(r["source"] or "").split("|")[0].strip() for r in recs}),
        "is_pushback_ai": dict(collections.Counter(m["is_pushback_ai"] for m in meta.values())),
        "team_notes_filled": sum(1 for m in meta.values() if m.get("team_notes") not in (None, "")),
        "subtheme_counts_ai": dict(sorted(sub_counts.items())),
        "theme_counts_ai": dict(theme_counts),
        "messages_per_case": {
            "min": min(r["n_messages"] for r in recs),
            "median": statistics.median(r["n_messages"] for r in recs),
            "max": max(r["n_messages"] for r in recs),
        },
        "target_chars": {"median": statistics.median(tgt_len), "p90": sorted(tgt_len)[int(.9*len(tgt_len))], "max": max(tgt_len)},
        "preceding_answer_chars": {"median": statistics.median(ans_len), "p90": sorted(ans_len)[int(.9*len(ans_len))], "max": max(ans_len)},
        "context_chars_before_target": {"median": statistics.median(ctx_len), "p90": sorted(ctx_len)[int(.9*len(ctx_len))], "max": max(ctx_len)},
        "target_is_last_user_turn": len(recs) - len(bad_target),
        "has_preceding_answer": len(recs) - len(no_prior),
    }
    with open(os.path.join(RESULTS, "corpus_stats.json"), "w") as f:
        json.dump(stats, f, indent=2)
    print(json.dumps(stats, indent=2)[:1200])
    print(f"\n[written {DATA_LOCAL}/*.jsonl (local, do not commit), {RESULTS}/corpus_stats.json]")


if __name__ == "__main__":
    main()
