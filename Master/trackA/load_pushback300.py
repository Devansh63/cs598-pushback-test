"""Track A: parse the pushback_300 materials (Devansh/progress/phase1/pushback300, ported).

Inputs (NOT in the repo; third-party conversation text):
  --contexts  pushback_300_contexts.md   full conversation context for 300 candidate pushbacks
  --workbook  pushback_300_coded.xlsm    300 rows of metadata + AI preliminary codes (read-only, no macros)

Outputs:
  results/trackA/data_local/contexts.jsonl   event_id, messages, target_index   (gitignored)
  results/trackA/data_local/meta.jsonl       workbook metadata by event_id      (gitignored)
  results/trackA/corpus_stats.json           aggregate counts only              (committed)
"""
import argparse
import collections
import json
import re
import statistics
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "results" / "trackA"
DATA_LOCAL = OUT / "data_local"

BLOCK_RE = re.compile(r"^## (\d{3})\.\s+(PB_[0-9a-f]+)\s*$", re.M)
HEAD_RE = re.compile(r"^\*\*Message (\d+): (user|assistant)([^*]*)\*\*\s*$", re.M)
FENCE_RE = re.compile(r"```(?:text)?\n(.*?)\n```", re.S)
SOURCE_RE = re.compile(r"Source:\s*`([^`]+)`")
FLAG_RE = re.compile(r"AI is_pushback:\s*(\w+)", re.I)


def parse_contexts(path):
    text = Path(path).read_text(encoding="utf-8")
    starts = [(m.start(), m.group(1), m.group(2)) for m in BLOCK_RE.finditer(text)]
    recs = []
    for i, (pos, no, eid) in enumerate(starts):
        end = starts[i + 1][0] if i + 1 < len(starts) else len(text)
        block = text[pos:end]
        parts = HEAD_RE.split(block)
        messages = []
        for j in range(1, len(parts), 4):
            num, role, flag, rest = parts[j], parts[j + 1], parts[j + 2], parts[j + 3]
            fence = FENCE_RE.search(rest)
            messages.append({"index": int(num), "role": role.strip(),
                             "flag": flag.strip().strip("—").strip(),
                             "content": (fence.group(1) if fence else rest).strip()})
        if not messages:
            continue
        src = SOURCE_RE.search(block)
        target = [m for m in messages if "TARGET USER TURN" in m["flag"]]
        prior = [m for m in messages if "PRECEDING ASSISTANT ANSWER" in m["flag"]]
        recs.append({"sample_no": int(no), "event_id": eid, "source": src.group(1) if src else None,
                     "is_pushback_ai": FLAG_RE.search(block).group(1).lower() if FLAG_RE.search(block) else None,
                     "messages": messages, "target_index": target[0]["index"] if target else None,
                     "target_text": target[0]["content"] if target else None,
                     "preceding_answer": prior[-1]["content"] if prior else None, "n_messages": len(messages)})
    return recs


def parse_workbook(path):
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl is required: pip install openpyxl")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_vba=False)
    ws = wb["Pushback 300"] if "Pushback 300" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    g = lambda r, i: r[i] if i < len(r) else None
    start = next((i + 1 for i, r in enumerate(rows) if str(g(r, 0) or "").strip().lower().startswith("sample")), None)
    if start is None:
        sys.exit("could not find the header row in the workbook")
    meta = {}
    for r in rows[start:]:
        eid = g(r, 1)
        if not eid:
            continue
        split = lambda v: [t.strip() for t in re.split(r"[\n|]", str(v or "")) if t.strip()]
        meta[str(eid).strip()] = {"sample_no": g(r, 0), "prev_user_excerpt": g(r, 2), "answer_excerpt": g(r, 3),
                                  "is_pushback_ai": str(g(r, 4) or "").strip().lower(), "ai_rationale": g(r, 5),
                                  "themes_ai": split(g(r, 6)), "subthemes_ai": split(g(r, 7)), "quote": g(r, 8),
                                  "team_notes": g(r, 9), "source_ref": g(r, 10)}
    return meta


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--contexts", required=True)
    ap.add_argument("--workbook", required=True)
    a = ap.parse_args()
    DATA_LOCAL.mkdir(parents=True, exist_ok=True)
    recs, meta = parse_contexts(a.contexts), parse_workbook(a.workbook)
    have = sum(1 for r in recs if r["event_id"] in meta)
    bad_target = [r for r in recs if r["target_index"] is None or r["target_index"] != r["messages"][-1]["index"] or r["messages"][-1]["role"] != "user"]
    no_prior = [r for r in recs if not r["preceding_answer"]]
    with open(DATA_LOCAL / "contexts.jsonl", "w", encoding="utf-8") as f:
        for r in recs:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    with open(DATA_LOCAL / "meta.jsonl", "w", encoding="utf-8") as f:
        for eid, m in meta.items():
            f.write(json.dumps({"event_id": eid, **m}, ensure_ascii=False) + "\n")
    sub = collections.Counter(s.split()[0].strip("—-") for m in meta.values() for s in m["subthemes_ai"] if s.split())
    q = lambda xs, p: sorted(xs)[int(p * len(xs))] if xs else None
    tgt = [len(r["target_text"] or "") for r in recs]
    ctx = [sum(len(m["content"]) for m in r["messages"][:-1]) for r in recs]
    stats = {"n_cases": len(recs), "n_joined": have, "target_is_last_user_turn": len(recs) - len(bad_target),
             "has_preceding_answer": len(recs) - len(no_prior),
             "is_pushback_ai": dict(collections.Counter(m["is_pushback_ai"] for m in meta.values())),
             "team_notes_filled": sum(1 for m in meta.values() if m.get("team_notes") not in (None, "")),
             "subtheme_counts_ai": dict(sorted(sub.items())),
             "messages_per_case": {"min": min(r["n_messages"] for r in recs), "median": statistics.median(r["n_messages"] for r in recs), "max": max(r["n_messages"] for r in recs)},
             "target_chars": {"median": statistics.median(tgt), "p90": q(tgt, .9), "max": max(tgt)},
             "context_chars_before_target": {"median": statistics.median(ctx), "p90": q(ctx, .9), "max": max(ctx)}}
    (OUT / "corpus_stats.json").write_text(json.dumps(stats, indent=2))
    print(json.dumps(stats, indent=2))


if __name__ == "__main__":
    main()
